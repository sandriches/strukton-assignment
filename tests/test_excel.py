from pathlib import Path

import openpyxl
import pytest

from eisen.excel import PROPOSAL_COLUMNS, REVIEW_COLUMNS, RUN_COLUMNS, DocumentResult, write_workbook
from eisen.extraction import Extraction
from eisen.pdf_text import Document
from eisen.verification import Proposal

DOC = Document(name="01_test.pdf", pages=["tekst", ""], removed=["Pagina 1", "Pagina 1"])


def proposal(text, **overrides):
    values = dict(
        document="01_test.pdf", code="VS-1", code_source="bron", text=text, pages="1",
        section="1.1", type="Bindend", model_reason="Reden.", doubt=False, found=True,
    )
    return Proposal(**{**values, **overrides})


@pytest.fixture
def workbook(tmp_path):
    proposals = [
        proposal("De Opdrachtnemer dient alles  op te leveren.", signals=["Signaal een", "Signaal twee"]),
        proposal("=SUM(A1:A2) is geen formule", code="VS-2"),
        proposal(
            "Gecorrigeerde tekst", code="D01-1.1-01", code_source="afgeleid", pages="", found=False,
            closest="Gecorigeerde tekst.", review_reason="Eistekst niet letterlijk in het document gevonden",
        ),
    ]
    extraction = Extraction(
        candidates=[],
        meta={
            "model_requested": "claude-sonnet-5-5", "model_served": "claude-sonnet-5-5",
            "effort": "high", "prompt": "extract_v1.md", "prompt_sha256": "abcdef1234567890",
            "created_at": "2026-10-28T10:00:00+00:00", "input_tokens": 100, "output_tokens": 50,
        },
        path=Path("output/responses/01_test.abc.json"),
        from_cache=True,
    )
    path = tmp_path / "eisen.xlsx"
    write_workbook(path, [DocumentResult(DOC, extraction, proposals)])
    return openpyxl.load_workbook(path)


def rows(sheet):
    return [list(r) for r in sheet.iter_rows(min_row=2, values_only=True)]


def test_sheets_and_headers(workbook):
    assert workbook.sheetnames == ["Voorstellen", "Te controleren", "Run-info"]
    for name, columns in zip(workbook.sheetnames, (PROPOSAL_COLUMNS, REVIEW_COLUMNS, RUN_COLUMNS)):
        assert [c.value for c in workbook[name][1]] == [header for header, _ in columns]


def test_verified_proposals_are_on_the_main_sheet_verbatim(workbook):
    main = rows(workbook["Voorstellen"])
    assert len(main) == 2
    assert main[0][:7] == [
        "01_test.pdf", "VS-1", "bron", "De Opdrachtnemer dient alles  op te leveren.", "1", "1.1", "Bindend",
    ]
    assert main[0][7] == "Signaal een\nSignaal twee"
    assert main[0][9] is None  # review column is left empty for the system engineer


def test_text_starting_with_equals_sign_stays_text(workbook):
    cell = workbook["Voorstellen"]["D3"]
    assert cell.value == "=SUM(A1:A2) is geen formule"
    assert cell.data_type == "s"


def test_review_sheet_shows_reason_and_closest_source_text(workbook):
    [row] = rows(workbook["Te controleren"])
    assert row[:7] == [
        "01_test.pdf", "D01-1.1-01", "afgeleid", "Eistekst niet letterlijk in het document gevonden",
        "Gecorrigeerde tekst", "nee", "Gecorigeerde tekst.",
    ]


def test_review_column_has_a_dropdown_on_both_sheets(workbook):
    for name, column in (("Voorstellen", "J"), ("Te controleren", "M")):
        [validation] = workbook[name].data_validations.dataValidation
        assert validation.formula1 == '"Goedgekeurd,Afgekeurd"'
        assert str(validation.sqref) == f"{column}2:{column}1000"


def test_run_info_records_model_prompt_and_verbatim_rate(workbook):
    run = workbook["Run-info"]
    row = dict(zip([c.value for c in run[1]], [c.value for c in run[2]]))
    assert row["Model (antwoord)"] == "claude-sonnet-5-5"
    assert row["Prompt"] == "extract_v1.md"
    assert row["Voorstellen totaal"] == 3
    assert row["Letterlijk gevonden"] == 2
    assert row["Letterlijk %"] == 66.7
    assert (row["Op Voorstellen"], row["Te controleren"]) == (2, 1)
    assert row["Verwijderde kop-/voetregels"] == "Pagina 1"
    assert row["Pagina's zonder tekstlaag"] == "2"
    assert run.max_row == 2  # no generation timestamp: same responses give the same workbook
