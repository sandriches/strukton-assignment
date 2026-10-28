"""Score the Excel output against the reference list.

Follows the scoring rules of the assignment ("Meetwijze"):

- A reference item is found when a proposal on the main sheet has the same text
  (whitespace and Unicode form ignored, as in source_index.normalize) and the
  same document and page (range). Section, code and type do not count.
- Several proposals together count as one found item when they cover its full
  text in order, without gaps or additions, and their combined page range
  equals the item's. Such splits are reported separately.
- Every other proposal on the main sheet is a duplicate (it matches an item that
  is already found) or wrong (onterecht), with the reason.
- The review sheet is scored separately: which of its rows are reference items.

The evaluation reads the workbook, not the model responses: it scores what the
system engineer receives, and needs no model connection. The report is read by
the user, so it is in Dutch.
"""

import difflib
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from eisen.excel import MAIN_SHEET, REVIEW_SHEET, RUN_SHEET
from eisen.source_index import normalize

FOUND = "gevonden"
SPLIT = "splitsing"
DUPLICATE = "dubbel"
WRONG = "onterecht"
REFERENCE_ITEM = "referentie-item"
OTHER = "geen referentie-item"


@dataclass(frozen=True)
class Reference:
    code: str
    text: str
    type: str
    pages: str


@dataclass(frozen=True)
class Row:
    sheet: str
    number: int  # Excel row number, so the user can find it back
    document: str
    code: str
    text: str
    pages: str
    type: str


@dataclass(frozen=True)
class Outcome:
    row: Row
    status: str
    reference: Reference | None = None
    reason: str = ""


@dataclass
class Score:
    document: str
    references: list[Reference]
    proposals: list[Outcome] = field(default_factory=list)  # main sheet, in sheet order
    review: list[Outcome] = field(default_factory=list)  # review sheet, in sheet order
    missed: list[tuple[Reference, Row | None]] = field(default_factory=list)  # with the closest row
    splits: dict[Reference, list[Row]] = field(default_factory=dict)  # parts in text order

    @property
    def found(self) -> int:
        return len(self.references) - len(self.missed)

    @property
    def rate(self) -> float | None:
        return 100 * self.found / len(self.references) if self.references else None

    def count(self, status: str) -> int:
        return sum(o.status == status for o in self.proposals)

    @property
    def type_mismatches(self) -> list[Outcome]:
        return [o for o in self.proposals if o.status in (FOUND, SPLIT) and o.row.type != o.reference.type]

    @property
    def review_references(self) -> int:
        return sum(o.status == REFERENCE_ITEM for o in self.review)


def load_reference(path: Path) -> dict[int, list[Reference]]:
    """Reference items per document number, from sheets named "Doc <n> ..."."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    references = {}
    for sheet in workbook.worksheets:
        number = re.match(r"Doc\s*(\d+)", sheet.title)
        if not number:
            continue
        records = _records(sheet, header="Eiscode")
        references[int(number[1])] = [
            Reference(
                code=_text(r["Eiscode"]),
                text=_text(r["Eistekst (letterlijk)"]),
                type=_text(r["Type"]),
                pages=_text(r["Pagina"]),
            )
            for _, r in records
            if r["Eistekst (letterlijk)"]
        ]
    return references


def load_output(path: Path) -> tuple[list[str], list[Row]]:
    """The documents in the workbook (from Run-info) and the rows of both proposal sheets."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    documents = [_text(r["Document"]) for _, r in _records(workbook[RUN_SHEET], header="Document")]
    rows = [
        Row(
            sheet=name,
            number=number,
            document=_text(r["Document"]),
            code=_text(r["Eiscode"]),
            text=_text(r["Eistekst"]),
            pages=_text(r["Pagina"]),
            type=_text(r["Type"]),
        )
        for name in (MAIN_SHEET, REVIEW_SHEET)
        for number, r in _records(workbook[name], header="Document")
    ]
    return [d for d in documents if d], rows


def document_number(name: str) -> int | None:
    """01_Vraagspecificatie.pdf → 1, matching the reference sheet "Doc 1 ..."."""
    number = re.match(r"(\d+)", name)
    return int(number[1]) if number else None


def score(document: str, references: list[Reference], rows: list[Row]) -> Score:
    rows = [r for r in rows if r.document == document]
    main = [r for r in rows if r.sheet == MAIN_SHEET]
    result = Score(document, references)

    found: dict[Reference, list[Row]] = {}
    outcomes: dict[Row, Outcome] = {}
    for row in main:
        hits = [ref for ref in references if _matches(row, ref)]
        new = [ref for ref in hits if ref not in found]
        if new:
            found[new[0]] = [row]
            outcomes[row] = Outcome(row, FOUND, new[0])
        elif hits:
            first = found[hits[0]][0]
            reason = f"Zelfde eis als {hits[0].code}, al gevonden in rij {first.number}"
            outcomes[row] = Outcome(row, DUPLICATE, hits[0], reason)

    for ref in references:
        if ref in found:
            continue
        parts = _split(ref, [r for r in main if r not in outcomes])
        if parts:
            found[ref] = result.splits[ref] = parts
            for i, part in enumerate(parts, 1):
                outcomes[part] = Outcome(part, SPLIT, ref, f"Deel {i} van {len(parts)} van {ref.code}")

    for row in main:
        if row not in outcomes:
            outcomes[row] = Outcome(row, WRONG, reason=_wrong_reason(row, references))
    result.proposals = [outcomes[row] for row in main]

    for row in rows:
        if row.sheet != REVIEW_SHEET:
            continue
        ref = next((ref for ref in references if _matches(row, ref)), None)
        if ref:
            where = "ook gevonden op Voorstellen" if ref in found else "ontbreekt op Voorstellen"
            result.review.append(Outcome(row, REFERENCE_ITEM, ref, f"{ref.code} ({where})"))
        else:
            result.review.append(Outcome(row, OTHER, reason=_wrong_reason(row, references)))

    result.missed = [(ref, _closest(ref, rows)) for ref in references if ref not in found]
    return result


def _matches(row: Row, ref: Reference) -> bool:
    return normalize(row.text) == normalize(ref.text) and _same_pages(row.pages, ref.pages)


def _split(ref: Reference, rows: list[Row]) -> list[Row] | None:
    """Rows that together cover the reference text exactly, in order; None if there are none."""
    target = normalize(ref.text)
    pieces = [(normalize(r.text), r) for r in rows]
    pieces = [(text, r) for text, r in pieces if text and text != target and text in target]

    def cover(start: int, used: list[Row]) -> list[Row] | None:
        if start == len(target):
            return used
        for text, row in pieces:
            if row not in used and target.startswith(text, start):
                parts = cover(start + len(text), [*used, row])
                if parts:
                    return parts
        return None

    parts = cover(0, [])
    if parts and _page_range([p.pages for p in parts]) == _pages(ref.pages):
        return parts
    return None


def _wrong_reason(row: Row, references: list[Reference]) -> str:
    text = normalize(row.text)
    for ref in references:
        if text == normalize(ref.text):
            return f"Tekst klopt met {ref.code}, maar pagina {row.pages or '–'} in plaats van {ref.pages}"
    for ref in references:
        if text and text in normalize(ref.text):
            return f"Onvolledig: deel van {ref.code}"
    for ref in references:
        if normalize(ref.text) in text:
            return f"Bevat {ref.code} met extra tekst"
    return "Geen eis volgens de referentielijst"


def _closest(ref: Reference, rows: list[Row]) -> Row | None:
    target = normalize(ref.text)
    ranked = [(difflib.SequenceMatcher(None, target, normalize(r.text)).ratio(), r) for r in rows]
    ranked = [(ratio, r) for ratio, r in ranked if ratio >= 0.5]
    return max(ranked, key=lambda pair: pair[0])[1] if ranked else None


def _pages(pages: str) -> tuple[int, int] | None:
    """"2" → (2, 2), "1-2" → (1, 2); None for anything else, such as an empty cell."""
    numbers = re.fullmatch(r"\s*(\d+)\s*(?:[-–]\s*(\d+)\s*)?", pages)
    if not numbers:
        return None
    start = int(numbers[1])
    return start, int(numbers[2] or start)


def _same_pages(a: str, b: str) -> bool:
    return _pages(a) is not None and _pages(a) == _pages(b)


def _page_range(pages: list[str]) -> tuple[int, int] | None:
    ranges = [_pages(p) for p in pages]
    if not ranges or None in ranges:
        return None
    return min(start for start, _ in ranges), max(end for _, end in ranges)


def _records(sheet, header: str):
    """(row number, {column header: value}) for each non-empty row below the header row."""
    columns = None
    for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        if columns is None:
            if header in values:
                columns = values
            continue
        if any(v not in (None, "") for v in values):
            yield number, dict(zip(columns, values))


def _text(value) -> str:
    return "" if value is None else str(value).strip()


# Report


def write_report(path: Path, scores: list[Score], workbook: Path, reference: Path, skipped: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_report(scores, workbook, reference, skipped), encoding="utf-8")


def summary_line(s: Score) -> str:
    return (
        f"{s.document}: {s.found}/{len(s.references)} gevonden ({_percent(s.rate)}), "
        f"{s.count(WRONG)} onterecht, {s.count(DUPLICATE)} dubbel, {len(s.splits)} gesplitst; "
        f"Te controleren: {len(s.review)}, waarvan {s.review_references} referentie-item"
    )


def render_report(scores: list[Score], workbook: Path, reference: Path, skipped: list[str]) -> str:
    lines = [
        "# Evaluatie tegen de referentielijst",
        "",
        f"Uitvoer: `{workbook}`. Referentielijst: `{reference}`.",
        "",
        "Een referentie-item telt als gevonden als een voorstel op *Voorstellen* dezelfde tekst heeft "
        "(spaties, regeleinden en Unicode-vorm genegeerd) en hetzelfde document en dezelfde pagina of "
        "hetzelfde paginabereik. Paragraaf, eiscode en type tellen niet mee. Meerdere voorstellen tellen "
        "samen als één gevonden item als ze de volledige tekst op volgorde dekken, niets toevoegen en "
        "samen hetzelfde paginabereik hebben (splitsing). *Te controleren* telt niet mee in het "
        "vindpercentage en wordt apart gerapporteerd.",
        "",
        "## Samenvatting",
        "",
        "| Document | Gevonden | Vindpercentage | Onterecht | Dubbel | Gesplitst | Type gelijk "
        "| Te controleren | waarvan referentie-item |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for s in scores:
        lines.append(
            f"| {s.document} | {s.found}/{len(s.references)} | {_percent(s.rate)} | {s.count(WRONG)} "
            f"| {s.count(DUPLICATE)} | {len(s.splits)} | {s.found - len(_mismatched_references(s))}/{s.found} "
            f"| {len(s.review)} | {s.review_references} |"
        )
    if skipped:
        lines += ["", "Niet geëvalueerd: " + "; ".join(skipped) + "."]

    for s in scores:
        lines += ["", f"## {s.document}"]
        lines += _section("Gemist", [_missed(ref, row) for ref, row in s.missed])
        lines += _section("Gesplitst", [
            f"**{ref.code}** (pagina {ref.pages}): rijen {', '.join(str(r.number) for r in rows)}"
            for ref, rows in s.splits.items()
        ])
        lines += _section("Onterecht", [_outcome(o) for o in s.proposals if o.status == WRONG])
        lines += _section("Dubbel", [_outcome(o) for o in s.proposals if o.status == DUPLICATE])
        lines += _section("Type wijkt af van de referentielijst", [
            f"Rij {o.row.number} ({o.reference.code}): {o.row.type} in plaats van {o.reference.type}"
            for o in s.type_mismatches
        ])
        lines += _section(f"Te controleren ({len(s.review)})", [_outcome(o) for o in s.review])
    return "\n".join(lines) + "\n"


def _section(title: str, items: list[str]) -> list[str]:
    return ["", f"### {title}", "", *([f"- {item}" for item in items] or ["Geen."])]


def _missed(ref: Reference, closest: Row | None) -> str:
    line = f"**{ref.code}** (pagina {ref.pages}): \"{ref.text}\""
    if closest:
        line += f" Dichtstbijzijnde voorstel: {closest.sheet} rij {closest.number} (pagina {closest.pages or '–'})."
    return line


def _outcome(o: Outcome) -> str:
    code = f", {o.row.code}" if o.row.code else ""
    return f"Rij {o.row.number}{code} (pagina {o.row.pages or '–'}): {o.reason or o.status}. \"{o.row.text}\""


def _mismatched_references(s: Score) -> set[Reference]:
    return {o.reference for o in s.type_mismatches}


def _percent(rate: float | None) -> str:
    return "–" if rate is None else f"{rate:.1f}%".replace(".", ",")
