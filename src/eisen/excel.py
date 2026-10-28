"""Write proposals to an Excel workbook for review by a system engineer.

Three sheets, with Dutch names and headers because the user reads them:

- Voorstellen: verified proposals, with an empty review column (dropdown).
- Te controleren: proposals that need a human decision, with the reason.
- Run-info: per document, which model and prompt produced the output and how
  many quotes were found verbatim; the basis for spotting quality drops.
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from eisen.extraction import Extraction
from eisen.pdf_text import Document
from eisen.verification import Proposal

MAIN_SHEET = "Voorstellen"
REVIEW_SHEET = "Te controleren"
RUN_SHEET = "Run-info"

REVIEW_CHOICES = ("Goedgekeurd", "Afgekeurd")
REVIEW_ROWS = 1000  # rows that get the review dropdown, room for manual additions

# (header, width); wide columns wrap their text.
PROPOSAL_COLUMNS = [
    ("Document", 28),
    ("Eiscode", 14),
    ("Codebron", 10),
    ("Eistekst", 70),
    ("Pagina", 8),
    ("Paragraaf", 10),
    ("Type", 9),
    ("Signalering", 40),
    ("Toelichting model", 45),
    ("Beoordeling SE", 16),
]
REVIEW_COLUMNS = [
    ("Document", 28),
    ("Eiscode", 14),
    ("Codebron", 10),
    ("Reden", 30),
    ("Eistekst", 70),
    ("Letterlijk gevonden", 11),
    ("Dichtstbijzijnde brontekst", 70),
    ("Pagina", 8),
    ("Paragraaf", 10),
    ("Type", 9),
    ("Signalering", 40),
    ("Toelichting model", 45),
    ("Beoordeling SE", 16),
]
RUN_COLUMNS = [
    ("Document", 28),
    ("Model (gevraagd)", 20),
    ("Model (antwoord)", 20),
    ("Effort", 8),
    ("Prompt", 16),
    ("Prompt SHA-256", 18),
    ("Antwoord gemaakt op", 22),
    ("Modelantwoord", 50),
    ("Voorstellen totaal", 10),
    ("Letterlijk gevonden", 10),
    ("Letterlijk %", 10),
    ("Op Voorstellen", 10),
    ("Te controleren", 10),
    ("Tokens in", 10),
    ("Tokens uit", 10),
    ("Verwijderde kop-/voetregels", 40),
    ("Pagina's zonder tekstlaag", 12),
]

WRAP_FROM_WIDTH = 30
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")


@dataclass(frozen=True)
class DocumentResult:
    document: Document
    extraction: Extraction
    proposals: list[Proposal]


def write_workbook(path: Path, results: list[DocumentResult]) -> None:
    workbook = Workbook()
    main = workbook.active
    main.title = MAIN_SHEET
    review = workbook.create_sheet(REVIEW_SHEET)
    run = workbook.create_sheet(RUN_SHEET)

    _header(main, PROPOSAL_COLUMNS)
    _header(review, REVIEW_COLUMNS)
    _header(run, RUN_COLUMNS)

    for result in results:
        for p in result.proposals:
            signals = "\n".join(p.signals)
            if p.needs_review:
                _append(review, [
                    p.document, p.code, p.code_source, p.review_reason, p.text,
                    "ja" if p.found else "nee", p.closest or "", p.pages, p.section, p.type,
                    signals, p.model_reason, "",
                ])
            else:
                _append(main, [
                    p.document, p.code, p.code_source, p.text, p.pages, p.section, p.type,
                    signals, p.model_reason, "",
                ])
        _append(run, _run_row(result))

    for sheet, columns in ((main, PROPOSAL_COLUMNS), (review, REVIEW_COLUMNS)):
        _review_dropdown(sheet, len(columns))
    for sheet, columns in ((main, PROPOSAL_COLUMNS), (review, REVIEW_COLUMNS), (run, RUN_COLUMNS)):
        _layout(sheet, columns)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _run_row(result: DocumentResult) -> list:
    meta, proposals = result.extraction.meta, result.proposals
    found = sum(p.found for p in proposals)
    review = sum(p.needs_review for p in proposals)
    return [
        result.document.name,
        meta.get("model_requested", ""),
        meta.get("model_served", ""),
        meta.get("effort", ""),
        meta.get("prompt", ""),
        meta.get("prompt_sha256", "")[:12],
        meta.get("created_at", ""),
        str(result.extraction.path),
        len(proposals),
        found,
        round(100 * found / len(proposals), 1) if proposals else 0.0,
        len(proposals) - review,
        review,
        meta.get("input_tokens", ""),
        meta.get("output_tokens", ""),
        "\n".join(sorted(set(result.document.removed))),
        ", ".join(map(str, result.document.pages_without_text)),
    ]


def _header(sheet: Worksheet, columns: list[tuple[str, int]]) -> None:
    sheet.append([name for name, _ in columns])
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _append(sheet: Worksheet, values: list) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        # Requirement text is data, never a formula: "=..." must stay literal text.
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _review_dropdown(sheet: Worksheet, column: int) -> None:
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1=f'"{",".join(REVIEW_CHOICES)}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Ongeldige beoordeling",
        error=f"Kies {' of '.join(REVIEW_CHOICES)}.",
    )
    validation.add(f"{letter}2:{letter}{REVIEW_ROWS}")
    sheet.add_data_validation(validation)


def _layout(sheet: Worksheet, columns: list[tuple[str, int]]) -> None:
    for i, (_, width) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell, (_, width) in zip(row, columns):
            cell.alignment = Alignment(wrap_text=width >= WRAP_FROM_WIDTH, vertical="top")
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
